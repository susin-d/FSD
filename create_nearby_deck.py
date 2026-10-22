from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pathlib import Path

OUT = Path(__file__).parent / 'deliverables' / 'Nearby_Project_Presentation.pptx'
OUT.parent.mkdir(parents=True, exist_ok=True)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'Nearby | Local Services Marketplace'
prs.core_properties.subject = 'Project overview, requirements, architecture, implementation and roadmap'
prs.core_properties.author = 'Nearby Project Team'

SW, SH = 13.333, 7.5
C = {
    'ink':'15191D','muted':'625F54','blue':'1557FF','blue2':'DCE7FF','yellow':'FFDA35',
    'cream':'FFFDF7','white':'FFFFFF','line':'E5E0D6','green':'2B7852','green2':'E3F1E7',
    'orange':'E9784E','orange2':'FBE9DF','slate':'2B343E','light':'F5F3EC','red':'B64638'
}
def rgb(h): return RGBColor.from_string(C[h] if h in C else h)
def rect(slide,x,y,w,h,fill, radius=False, line=None, lw=1):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    sh.line.color.rgb=rgb(line or fill); sh.line.width=Pt(lw)
    if radius:
        try: sh.adjustments[0]=0.16
        except: pass
    return sh
def textbox(slide,x,y,w,h,text,size=16,color='ink',bold=False,font='Aptos',align=PP_ALIGN.LEFT,valign=MSO_ANCHOR.TOP,margin=0.03):
    box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=box.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=valign
    for i,line in enumerate(str(text).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.alignment=align
        p.space_after=Pt(5); p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
    return box
def dot(slide,x,y,d,fill):
    s=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); s.fill.solid(); s.fill.fore_color.rgb=rgb(fill); s.line.color.rgb=rgb(fill); return s
def line(slide,x1,y1,x2,y2,color='line',width=1.5,dash=None):
    s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); s.line.color.rgb=rgb(color); s.line.width=Pt(width)
    if dash: s.line.dash_style=dash
    return s
def pill(slide,x,y,w,label,fill='blue2',color='blue'):
    rect(slide,x,y,w,.32,fill,True)
    textbox(slide,x+.06,y+.015,w-.12,.27,label,10,color,True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
def base(title,kicker=None,dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg='slate' if dark else 'cream'; s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(bg)
    if not dark:
        rect(s,0,0,.13,SH,'yellow')
        if kicker: textbox(s,.65,.36,6,.28,kicker.upper(),10,'blue',True)
        textbox(s,.65,.72,12.0,.62,title,29,'ink',True)
        line(s,.65,1.47,12.68,1.47,'line',1)
        textbox(s,.65,7.12,5,.2,'NEARBY  /  PROJECT PRESENTATION',8,'muted',True)
        textbox(s,12.0,7.08,.65,.24,f'{len(prs.slides):02}',9,'muted',True,align=PP_ALIGN.RIGHT)
    return s
def card(slide,x,y,w,h,title,body,accent='blue',body_size=14):
    rect(slide,x,y,w,h,'white',True,'line',.75); rect(slide,x,y,.07,h,accent,True)
    textbox(slide,x+.22,y+.18,w-.4,.38,title,16,'ink',True)
    textbox(slide,x+.22,y+.68,w-.4,h-.78,body,body_size,'muted')
def bullet_list(slide,x,y,w,items,size=15,gap=.54,color='ink',dot_color='blue'):
    for i,item in enumerate(items):
        yy=y+i*gap; dot(slide,x,yy+.12,.09,dot_color); textbox(slide,x+.2,yy,w-.22,gap,item,size,color)
def node(slide,x,y,w,h,title,sub,fill='white',edge='line',titlecolor='ink',subcolor='muted'):
    rect(slide,x,y,w,h,fill,True,edge,1.2); textbox(slide,x+.09,y+.12,w-.18,.33,title,15,titlecolor,True,align=PP_ALIGN.CENTER); textbox(slide,x+.12,y+.52,w-.24,h-.58,sub,10.5,subcolor,align=PP_ALIGN.CENTER)

# 1 Title
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('cream')
rect(s,0,0,7.2,SH,'yellow'); textbox(s,.8,.75,5.6,.32,'LOCAL SERVICES MARKETPLACE',11,'blue',True)
textbox(s,.75,1.35,6.1,1.25,'Nearby',56,'ink',True,font='Aptos Display')
textbox(s,.82,2.8,5.5,1.15,'Find trusted services\nright where you live.',25,'ink',False)
line(s,.82,4.35,5.7,4.35,'ink',1)
textbox(s,.82,4.62,5.5,.48,'Project overview  /  System design  /  Implementation',14,'ink',True)
textbox(s,.82,6.82,5,.22,'WEB PLATFORM  •  INDIA',9,'ink',True)
# Decorative service marketplace map-like composition
rect(s,7.2,0,6.133,7.5,'cream')
for x,y,w,h,col in [(8.2,.75,4.2,1.2,'white'),(7.8,2.2,3.2,1.4,'blue'),(10.8,2.35,1.8,1.25,'white'),(8.5,4.1,4.3,1.15,'white'),(7.75,5.55,2.2,1.05,'orange2'),(10.2,5.55,2.65,1.05,'green2')]: rect(s,x,y,w,h,col,True)
dot(s,8.55,1.1,.48,'yellow'); textbox(s,9.18,1.05,2.9,.55,'Services nearby',18,'ink',True,valign=MSO_ANCHOR.MIDDLE)
textbox(s,8.15,2.5,2.55,.35,'PINCODE  560001',10,'white',True)
textbox(s,8.15,2.9,2.55,.52,'Exact-area\ndiscovery',18,'white',True)
textbox(s,11.0,2.65,1.45,.55,'Book a\nprovider',14,'ink',True,align=PP_ALIGN.CENTER)
textbox(s,8.78,4.4,3.7,.48,'A clear path from search to service',15,'ink',True,align=PP_ALIGN.CENTER)
for i,(lab,col) in enumerate([('DISCOVER','blue'),('BOOK','orange'),('COMPLETE','green')]):
    x=7.95+i*1.72; dot(s,x,5.86,.28,col); textbox(s,x+.36,5.82,1.22,.35,lab,9,'ink',True,valign=MSO_ANCHOR.MIDDLE)
textbox(s,7.75,6.85,4.9,.22,'CUSTOMER  /  PROVIDER  /  ADMIN',9,'muted',True)

# 2 Problem
s=base('Local service discovery is fragmented','01  /  problem')
textbox(s,.75,1.8,5.4,.55,'Customers and providers lack one reliable workflow.',22,'ink',True)
bullet_list(s,.82,2.65,5.3,['Customers search across disconnected listings and chats.','Service coverage and availability are hard to confirm.','Booking, payment, and issue resolution happen in separate steps.'],15,.82)
rect(s,6.65,1.85,5.95,4.6,'slate',True)
textbox(s,7.05,2.18,4.8,.35,'THE USER JOURNEY TODAY',11,'yellow',True)
for i,(lab,sub) in enumerate([('Search','Find a local listing'),('Verify','Check area and timing'),('Coordinate','Arrange the booking'),('Resolve','Handle payment or issue')]):
    yy=2.8+i*.78; dot(s,7.1,yy+.04,.3,'yellow'); textbox(s,7.58,yy,1.45,.3,lab,15,'white',True); textbox(s,9.3,yy,2.8,.34,sub,13,'white')
    if i<3: line(s,7.25,yy+.36,7.25,yy+.74,'white',1)
textbox(s,7.05,6.03,4.95,.3,'Nearby brings these steps into one marketplace.',13,'yellow',True)

# 3 Existing vs proposed
s=base('A structured marketplace replaces fragmented coordination','02  /  existing vs proposed')
textbox(s,.9,1.78,5.3,.36,'EXISTING EXPERIENCE',12,'muted',True)
textbox(s,7.0,1.78,5.3,.36,'PROPOSED WITH NEARBY',12,'blue',True)
line(s,6.65,1.82,6.65,6.55,'line',1)
bullet_list(s,.9,2.35,5.2,['Listings scattered across channels','Area and slot checks by manual messages','Booking status unclear to both sides','Payment handled outside service flow','No shared dispute or review record'],15,.7,'muted','orange')
bullet_list(s,7.0,2.35,5.25,['Exact six-digit pincode discovery','Provider availability and booking states','Customer and provider workspaces','Checkout after provider acceptance','Booking-linked chat, reviews, and disputes'],15,.7,'ink','blue')
pill(s,7.0,6.18,2.1,'ONE BOOKING RECORD','green2','green')

# 4 requirements
s=base('Requirements cover marketplace behavior and service quality','03  /  requirements')
card(s,.8,1.82,5.6,4.75,'Functional requirements','• Register and sign in as customer or provider\n• Create and manage provider profiles and services\n• Search by category and exact pincode\n• Publish availability and request bookings\n• Accept, reject, pay, cancel, and complete bookings\n• Message booking participants and submit reviews\n• Administer verification, disputes, settings, and audit', 'blue',14)
card(s,6.9,1.82,5.6,4.75,'Non-functional requirements','SECURITY  /  JWT auth, password hashing, role checks\n\nVALIDATION  /  schema validation, rate limits, CORS\n\nRELIABILITY  /  database-first startup, health checks\n\nPRIVACY  /  KYC documents remain with vendor\n\nUSABILITY  /  responsive English-language web app\n\nOPERATIONS  /  audit records for payment/admin actions','yellow',14)

# 5 architecture
s=base('Three tiers separate user experience, rules, and records','04  /  system architecture')
node(s,.85,2.0,3.1,2.05,'PRESENTATION TIER','React + Vite\nReact Router\nCustomer, provider, admin views','blue2','blue')
node(s,5.1,2.0,3.1,2.05,'APPLICATION TIER','Express API /api/v1\nJWT + role middleware\nSocket.IO + lifecycle worker','orange2','orange')
node(s,9.35,2.0,3.1,2.05,'DATA TIER','MongoDB + Mongoose\nUsers, services, bookings\nPayments, messages, audit','green2','green')
line(s,3.95,3.02,5.05,3.02,'blue',2); line(s,8.2,3.02,9.3,3.02,'orange',2)
textbox(s,3.98,2.65,1.0,.25,'HTTPS / JSON',9,'muted',True,align=PP_ALIGN.CENTER)
textbox(s,8.23,2.65,1.0,.25,'MODELS',9,'muted',True,align=PP_ALIGN.CENTER)
node(s,5.1,4.75,3.1,1.15,'EXTERNAL SERVICES','Razorpay  /  Email  /  KYC vendor','white','line')
line(s,6.65,4.05,6.65,4.68,'orange',1.6)
textbox(s,.9,6.35,11.6,.32,'Checkout, live KYC, email delivery, and production readiness require configured external accounts.',12,'muted',False,align=PP_ALIGN.CENTER)

# 6 Stack
s=base('The stack uses established web and data tools','05  /  technology stack')
stack=[('WEB','React 19','Vite, React Router, Tailwind CSS v4'),('API','Node.js + Express 5','CommonJS modules, REST endpoints'),('REALTIME','Socket.IO','Booking participant chat'),('DATABASE','MongoDB + Mongoose','Document models and indexes'),('SECURITY','JWT, bcryptjs, Zod','Helmet, CORS allowlist, rate limits'),('INTEGRATIONS','Razorpay, Nodemailer','Checkout/Route and email, when configured')]
for i,(a,b,d) in enumerate(stack):
    x=.85+(i%3)*4.15; y=1.9+(i//3)*2.2
    rect(s,x,y,3.75,1.75,'white',True,'line',.8); pill(s,x+.2,y+.18,1.38,a,'blue2','blue')
    textbox(s,x+.22,y+.66,3.3,.34,b,17,'ink',True); textbox(s,x+.22,y+1.12,3.25,.45,d,12,'muted')

# 7 use case
s=base('Three actors use separate but connected capabilities','06  /  use case diagram')
rect(s,2.5,1.75,8.1,4.95,'white',True,'line',1)
textbox(s,2.8,1.9,2.5,.27,'NEARBY MARKETPLACE',10,'blue',True)
actors=[(.72,2.2,'CUSTOMER'),(.72,4.0,'PROVIDER'),(11.0,3.15,'ADMIN')]
for x,y,label in actors:
    dot(s,x+.34,y,.35,'yellow' if label!='ADMIN' else 'blue')
    textbox(s,x,y+.42,1.15,.3,label,10,'ink',True,align=PP_ALIGN.CENTER)
cases=[(3.0,2.35,'Search by pincode'),(5.75,2.35,'Request booking'),(8.5,2.35,'Pay and review'),(3.0,3.55,'Manage services'),(5.75,3.55,'Set availability'),(8.5,3.55,'Accept booking'),(3.0,4.8,'Booking chat'),(5.75,4.8,'Raise dispute'),(8.5,4.8,'Review providers')]
for x,y,t in cases:
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(2.0), Inches(.62)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb('blue2'); sh.line.color.rgb=rgb('blue')
    textbox(s,x+.08,y+.12,1.84,.34,t,10,'ink',True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
for x1,y1,x2,y2 in [(1.86,2.55,3,2.65),(1.86,2.6,5.75,2.65),(1.86,2.65,8.5,2.65),(1.86,4.1,3,3.85),(1.86,4.1,5.75,3.85),(1.86,4.1,8.5,3.85),(1.86,4.1,3,5.1),(1.86,4.1,5.75,5.1),(1.86,4.1,8.5,5.1),(11,3.5,10.5,3.85),(11,3.5,7.75,5.1)]: line(s,x1,y1,x2,y2,'line',.9)
textbox(s,3.0,6.3,7,.28,'Admin also reviews providers, resolves disputes, updates policy, and checks audit records.',11,'muted',align=PP_ALIGN.CENTER)

# 8 ER
s=base('Core records connect accounts, services, and bookings','07  /  entity relationship diagram')
entities={
 'User':(.8,2.05,2.35,1.35,'name · email · role'),
 'ProviderProfile':(4.0,1.78,2.55,1.35,'user · status · areas'),
 'Service':(8.0,1.78,2.35,1.35,'provider · pincode · price'),
 'Booking':(4.7,4.15,2.65,1.55,'customer · provider · service\nslot · status · pricePaise'),
 'Availability':(.85,4.65,2.45,1.25,'provider · date · windows'),
 'Payment':(8.35,4.15,2.35,1.35,'booking · state · amount'),
 'Message / Review / Dispute':(10.95,4.35,1.95,1.35,'booking · actor · content')}
for name,(x,y,w,h,fields) in entities.items():
    rect(s,x,y,w,h,'white',True,'line',1); rect(s,x,y,w,.38,'blue' if name in ['Booking','User'] else 'slate',True)
    textbox(s,x+.1,y+.045,w-.2,.28,name,12,'white',True,align=PP_ALIGN.CENTER)
    textbox(s,x+.12,y+.5,w-.24,h-.55,fields,10,'muted',align=PP_ALIGN.CENTER)
rels=[((3.15,2.7),(4,2.45),'1 : 1'),((6.55,2.4),(8,2.4),'1 : many'),((5.25,3.15),(5.8,4.1),'1 : many'),((8.8,3.15),(6.8,4.15),'1 : many'),((3.3,5.1),(4.65,4.9),'1 : many'),((7.4,4.85),(8.3,4.85),'1 : 1'),((7.35,5.45),(10.9,5.0),'1 : many')]
for (a,b,lbl) in rels:
    line(s,*a,*b,'blue',1.35); mx=(a[0]+b[0])/2; my=(a[1]+b[1])/2-.18; textbox(s,mx-.42,my,.85,.2,lbl,8,'blue',True,align=PP_ALIGN.CENTER)
textbox(s,.85,6.45,11.8,.28,'Simplified logical model. Payment, review, dispute, and message records reference their related booking.',10,'muted')

# 9 modules
s=base('Modules group the marketplace around its service lifecycle','08  /  module description')
mods=[('Identity','Registration, login, JWT, role access'),('Provider','Profile, service area, verification reference'),('Catalog','Listings, exact-pincode search, fixed INR price'),('Availability','Date-scoped windows and conflict checks'),('Bookings','Request, accept, reject, hold, cancel, complete'),('Payments','Checkout, webhook events, refunds, transfers'),('Engagement','Booking chat, reviews, email notifications'),('Operations','Disputes, settings, admin actions, audit')]
for i,(h,b) in enumerate(mods):
    x=.8+(i%4)*3.15; y=1.85+(i//4)*2.25
    textbox(s,x,y,2.75,.37,h,17,'ink',True); line(s,x,y+.48,x+2.7,y+.48,['blue','orange','green','yellow'][i%4],2)
    textbox(s,x,y+.68,2.73,.86,b,12,'muted')

# 10 roles
s=base('Each role enters a workspace suited to its responsibilities','09  /  role-based workspaces')
roles=[('CUSTOMER','Discover','Search services by category and exact pincode.','Book','Request service, pay after acceptance, message, review, or dispute.','blue'),('PROVIDER','Configure','Maintain profile, service areas, listings, availability.','Fulfil','Respond to requests, coordinate bookings, track completion.','orange'),('ADMIN','Oversee','Review provider status and manage marketplace settings.','Resolve','Work dispute queue, payment/payout operations, and audit.','green')]
for i,(role,a,b,c,d,col) in enumerate(roles):
    x=.8+i*4.15; rect(s,x,1.9,3.78,4.55,'white',True,'line',.8); rect(s,x,1.9,3.78,.65,col,True)
    textbox(s,x+.2,2.08,3.3,.3,role,13,'white' if col!='yellow' else 'ink',True)
    textbox(s,x+.25,2.95,3.2,.33,a,18,'ink',True); textbox(s,x+.25,3.4,3.15,.9,b,13,'muted')
    line(s,x+.25,4.5,x+3.45,4.5,'line',1); textbox(s,x+.25,4.77,3.2,.33,c,18,'ink',True); textbox(s,x+.25,5.22,3.15,.9,d,13,'muted')

# 11 frontend
s=base('The frontend guides users through discovery and booking','10  /  frontend development')
textbox(s,.85,1.88,5.15,.65,'React app with role-aware routing and a shared session context.',19,'ink',True)
bullet_list(s,.92,2.8,5.25,['Home and service search pages','Service details and booking entry','Login and customer/provider signup','Protected customer, provider, and admin paths','Responsive layouts with warm yellow and blue branding'],14,.67)
rect(s,7.0,1.9,5.1,4.5,'light',True)
textbox(s,7.38,2.22,4.3,.28,'SCREEN FLOW',10,'blue',True)
node(s,7.45,2.8,1.65,1.0,'Home','Search','white','line')
node(s,9.75,2.8,1.75,1.0,'Service','Details','blue2','blue')
node(s,8.65,4.45,1.8,1.0,'Workspace','Role protected','green2','green')
line(s,9.1,3.3,9.7,3.3,'blue',1.6); line(s,10.6,3.82,9.55,4.4,'blue',1.6)
textbox(s,7.4,5.85,4.1,.28,'Router redirects unauthorized roles to their workspace.',11,'muted')

# 12 Backend
s=base('The API centralizes authorization and booking rules','11  /  backend development')
node(s,.9,2.0,2.45,1.2,'REQUEST','JSON + JWT','blue2','blue')
node(s,4.15,2.0,2.45,1.2,'MIDDLEWARE','Auth + role + limits','orange2','orange')
node(s,7.4,2.0,2.45,1.2,'MODULE ROUTE','Validate + authorize','green2','green')
node(s,10.3,2.0,2.2,1.2,'MODEL','Mongoose','white','line')
line(s,3.38,2.6,4.1,2.6,'blue',1.7); line(s,6.63,2.6,7.35,2.6,'orange',1.7); line(s,9.88,2.6,10.25,2.6,'green',1.7)
card(s,1.05,4.0,5.25,1.75,'Lifecycle controls','Slot holds expire. Payment webhooks protect against duplicate events. Open disputes pause payout.','blue',13)
card(s,7.05,4.0,5.25,1.75,'Operational safeguards','Health checks require database reachability. Admin and payment actions leave audit records.','orange',13)
textbox(s,1.0,6.3,11.3,.28,'Express API version: /api/v1   |   Common response envelope: success / data or error',11,'muted',True,align=PP_ALIGN.CENTER)

# 13 DB
s=base('MongoDB stores domain records with explicit relationships','12  /  database integration')
bullet_list(s,.95,2.0,5.2,['Mongoose schemas define fields, enums, and references.','Indexes support email uniqueness and booking lookups.','Money is stored in integer paise.','Provider verification stores status and vendor reference.','Booking state ties discovery, payment, chat, and review together.'],14,.75)
rect(s,6.75,1.95,5.55,4.5,'slate',True)
textbox(s,7.1,2.27,4.8,.3,'BOOKING AS THE TRANSACTION HUB',11,'yellow',True)
dot(s,8.85,3.45,.75,'yellow'); textbox(s,9.72,3.55,1.9,.38,'BOOKING',18,'white',True,valign=MSO_ANCHOR.MIDDLE)
for x,y,t in [(7.2,3.1,'Customer'),(10.75,3.1,'Provider'),(7.2,4.6,'Service'),(10.75,4.6,'Payment'),(8.7,5.45,'Chat / Review / Dispute')]:
    rect(s,x,y,1.55,.55,'white',True); textbox(s,x+.06,y+.12,1.43,.25,t,10,'ink',True,align=PP_ALIGN.CENTER)
line(s,8.7,3.65,8.95,3.8,'yellow',1.2); line(s,10.1,3.8,10.75,3.55,'yellow',1.2); line(s,8.7,4.05,8.0,4.6,'yellow',1.2); line(s,10.0,4.0,10.75,4.8,'yellow',1.2); line(s,9.2,4.2,9.4,5.42,'yellow',1.2)

# 14 APIs
s=base('Versioned APIs expose focused operations by domain','13  /  api development')
groups=[('/auth','Register, login, current user'),('/services','Search, list, update, delete'),('/availability','Publish and check time windows'),('/bookings','Request, accept, pay, complete'),('/payments','Webhook, verify, refund'),('/messages','REST history + Socket.IO'),('/admin','Verification, disputes, audit')]
for i,(route,desc) in enumerate(groups):
    yy=1.78+i*.65; rect(s,.9,yy,2.4,.42,'blue2',True); textbox(s,1.03,yy+.05,2.15,.28,route,12,'blue',True)
    textbox(s,3.65,yy+.05,6.3,.3,desc,14,'ink'); textbox(s,10.4,yy+.06,1.7,.25,'/api/v1',10,'muted',True)
    if i<len(groups)-1: line(s,.9,yy+.53,12.0,yy+.53,'line',.6)
rect(s,9.65,1.85,2.5,1.1,'slate',True); textbox(s,9.88,2.05,2.1,.65,'Bearer JWT\nAuthenticated routes',12,'white',True,align=PP_ALIGN.CENTER)
textbox(s,9.65,3.25,2.55,1.25,'Socket.IO events\nbooking:join\nmessage:send / message:new',11,'muted')
textbox(s,.92,6.55,11.2,.25,'Payment webhooks validate the raw request signature and guard duplicate event processing.',11,'muted')

# 15 module implementation
s=base('Implementation follows a controlled service lifecycle','14  /  module-wise implementation')
steps=[('01','Discover','Active service + exact pincode'),('02','Schedule','Provider time window + conflict check'),('03','Request','Booking linked to customer/provider/service'),('04','Accept','Temporary slot hold starts'),('05','Checkout','Payment after provider acceptance'),('06','Complete','Review window or dispute resolution')]
for i,(n,h,d) in enumerate(steps):
    x=.82+i*2.08; dot(s,x+.65,2.12,.52,'blue' if i<4 else 'orange'); textbox(s,x+.65,2.22,.52,.25,n,11,'white',True,align=PP_ALIGN.CENTER)
    if i<5: line(s,x+1.18,2.38,x+2.04,2.38,'line',2)
    textbox(s,x,2.95,1.85,.34,h,16,'ink',True,align=PP_ALIGN.CENTER); textbox(s,x+.02,3.42,1.82,1.05,d,11,'muted',align=PP_ALIGN.CENTER)
rect(s,1.15,5.15,11.0,.9,'yellow',True)
textbox(s,1.42,5.42,10.45,.34,'Each transition checks role, booking state, availability, and payment or dispute conditions.',14,'ink',True,align=PP_ALIGN.CENTER)

# 16 challenges
s=base('External dependencies and state changes need careful handling','15  /  challenges and solutions')
rows=[('Pincode precision','Exact six-digit match for service and provider area.'),('Double booking','Slot holds, overlap checks, and unpaid-hold expiry.'),('Payment event retries','Webhook signature validation and duplicate event protection.'),('KYC privacy','Vendor retains documents; Nearby stores references and status.'),('Unsettled payout','Hold transfers during refunds, failures, or open disputes.'),('Launch readiness','External credentials, policy, email, hosting, and operations remain configuration work.')]
for i,(a,b) in enumerate(rows):
    yy=1.75+i*.74; textbox(s,.9,yy,2.55,.35,a,14,'ink',True); line(s,3.65,yy+.18,4.02,yy+.18,'blue',1.5); textbox(s,4.25,yy,7.8,.48,b,13,'muted')

# 17 conclusion / future / thank you (combine required ending topics into one final slide while preserving 17 slides? count has 17 requested topics so this is slide 17, include both conclusion/future and thank you)
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('slate')
rect(s,0,0,.14,SH,'yellow'); textbox(s,.82,.58,5,.28,'16  /  CONCLUSION & FUTURE ENHANCEMENTS',10,'yellow',True)
textbox(s,.82,1.1,7.5,.7,'A connected path from search to service.',30,'white',True)
textbox(s,.85,2.15,5.45,1.0,'Nearby connects local discovery with booking, payment, communication, and issue handling.',18,'white')
textbox(s,.85,3.55,5,.32,'NEXT STEPS',11,'yellow',True)
bullet_list(s,.9,4.05,5.5,['Select and configure a production KYC vendor.','Complete Razorpay approval and sandbox-to-live validation.','Add operational monitoring, backups, and privacy workflows.','Improve search insights and provider tools from user feedback.'],13,.55,'white','yellow')
line(s,7.0,1.1,7.0,6.5,'white',.8)
textbox(s,7.65,2.05,4.7,.36,'THANK YOU',12,'yellow',True)
textbox(s,7.55,2.65,5.0,1.05,'Nearby',48,'white',True,font='Aptos Display')
textbox(s,7.7,3.9,4.6,.7,'Local services,\nwithin reach.',22,'white')
rect(s,7.7,5.3,3.8,.55,'blue',True); textbox(s,7.85,5.42,3.5,.3,'CUSTOMER  /  PROVIDER  /  ADMIN',10,'white',True,align=PP_ALIGN.CENTER)
textbox(s,.82,7.05,5,.2,'NEARBY  /  PROJECT PRESENTATION',8,'white',True)
textbox(s,12.0,7.02,.65,.24,'17',9,'yellow',True,align=PP_ALIGN.RIGHT)

prs.save(OUT)
print(OUT)
